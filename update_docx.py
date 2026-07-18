#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Update Comparatie_preturi_deviz.docx with live-scraped prices.

- Match by `item` (row = Nr. poziției) and `store` (column).
- Replace ONLY the numeric price text inside the existing hyperlink, so the blue
  underlined link and its target (source product page) are preserved.
- Confirmed-live prices overwrite the cell value.
- A per-row note is appended to "Observații" showing, for that position, which
  store prices were confirmed live (✓) and which stay neconfirmate, with date.
- Header gets a "Prețuri verificate live: <data>" line.
- Column widths / table layout are left untouched (only run text changes).
"""
import sys, json, re
from collections import defaultdict
import docx
from docx.oxml.ns import qn
from docx.shared import RGBColor

DOCX_IN = "/home/user/classifier/Comparatie_preturi_deviz.docx"
DOCX_OUT = "/home/user/classifier/Comparatie_preturi_deviz_actualizat.docx"
PRICES = "/home/user/classifier/prices.json"
LIVE_DATE = "18 iulie 2026"

STORE_COL = {"Dedeman": 2, "Leroy Merlin": 3, "Romstal": 4, "eMAG": 5, "Bricodepot": 6}
OBS_COL = 8


def ro_format(v):
    v = float(v)
    if abs(v - round(v)) < 0.005:
        return f"{int(round(v)):,}".replace(",", ".")
    return f"{v:,.2f}".replace(",", "\x00").replace(".", ",").replace("\x00", ".")


def hl_text(hl):
    return ''.join(n.text or '' for n in hl.iter(qn('w:t')))


def set_hl_text(hl, text):
    runs = hl.findall(qn('w:r'))
    if not runs:
        return False
    first = runs[0]
    for extra in runs[1:]:
        hl.remove(extra)
    for child in list(first):
        if child.tag in (qn('w:t'), qn('w:br')):
            first.remove(child)
    t = first.makeelement(qn('w:t'), {qn('xml:space'): 'preserve'})
    t.text = text
    first.append(t)
    return True


def update_price_cell(cell, ro_price):
    """Replace the leading price line only if it is actually a price."""
    paras = cell.paragraphs
    if not paras:
        return False, ""
    hl0 = paras[0]._p.find(qn('w:hyperlink'))
    if hl0 is None:
        return False, ""
    t0 = hl_text(hl0)
    if not re.search(r'\d', t0):
        return False, t0            # e.g. "la cerere" / "componente" -> don't touch
    t1 = ""
    if len(paras) > 1:
        hl1 = paras[1]._p.find(qn('w:hyperlink'))
        t1 = hl_text(hl1) if hl1 is not None else ""
    if 'lei' in t0.lower():
        set_hl_text(hl0, f"{ro_price} lei")
    elif t1.strip().lower() == 'lei':
        set_hl_text(hl0, ro_price)
    else:
        set_hl_text(hl0, f"{ro_price} lei")
    return True, t0.strip()


def add_obs_note(cell, text, color=None):
    p = cell.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.size = docx.shared.Pt(7.5)
    if color:
        r.font.color.rgb = color


def main():
    rows = json.load(open(PRICES, encoding="utf-8"))
    by_item = defaultdict(list)
    for r in rows:
        if r["store"] in STORE_COL and r["item"] <= 29:
            by_item[r["item"]].append(r)

    d = docx.Document(DOCX_IN)
    t = d.tables[0]

    # header live-date line, inserted right after the existing intro block
    p_new = d.paragraphs[1].insert_paragraph_before()
    run = p_new.add_run(f"Prețuri verificate live: {LIVE_DATE} — sursă: paginile-produs "
                        f"(extragere automată prin JSON-LD/meta). Vezi coloana „Observații” "
                        f"pentru statusul per magazin.")
    run.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x6F, 0x3C)

    changed, notes_added = [], 0
    for item in sorted(by_item):
        recs = by_item[item]
        obs_cell = t.rows[item].cells[OBS_COL]
        live_tok, neconf_tok = [], []
        for r in recs:
            col = STORE_COL[r["store"]]
            cell = t.rows[item].cells[col]
            if r.get("confirmed") and r.get("price"):
                ro = ro_format(r["price"])
                ok, old = update_price_cell(cell, ro)
                changed.append((item, r["store"], old, ro, ok))
                live_tok.append(f"{r['store']} {ro} lei✓")
            else:
                # not confirmed live
                m = str(r.get("method", ""))
                if m.startswith("price_log") and r.get("price"):
                    neconf_tok.append(f"{r['store']} ~{ro_format(r['price'])} lei "
                                      f"(preț listat {m.split()[-1]}, neconf. live)")
                elif "blocat" in m or "403" in m:
                    neconf_tok.append(f"{r['store']} (anti-bot, neconf. live)")
                elif "eroare" in m or "Timeout" in m:
                    neconf_tok.append(f"{r['store']} (inaccesibil, neconf. live)")
                else:
                    neconf_tok.append(f"{r['store']} (neconf. live)")
        # build the per-row audit note
        parts = []
        if live_tok:
            parts.append("Live " + LIVE_DATE + ": " + "; ".join(live_tok) + ".")
        if neconf_tok:
            parts.append("Neconfirmat: " + "; ".join(neconf_tok) + ".")
        if parts:
            add_obs_note(obs_cell, " ".join(parts), color=RGBColor(0x33, 0x33, 0x33))
            notes_added += 1

    d.save(DOCX_OUT)
    ok_n = sum(1 for c in changed if c[4])
    print(f"Saved {DOCX_OUT}")
    print(f"Cells overwritten with live price: {ok_n}")
    print(f"Rows annotated in Observații: {notes_added}")
    print("\n--- price cell changes ---")
    for item, store, old, ro, ok in changed:
        flag = "" if ok else "  (SKIP: not a price cell, only noted in Obs)"
        print(f"  item {item:>2} {store:<12} {old!r:>28} -> {ro} lei{flag}")


if __name__ == "__main__":
    main()
