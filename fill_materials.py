#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Fill live unit prices into ListaResurse -> ListaResurse_preturi.xlsx.

Adds, next to the existing columns, an auditable trail: matched product,
source link, and a confidence flag, so every price can be verified. Rows that
are %/labour allowances are left unpriced. Saves incrementally."""
import sys, time, traceback
sys.path.insert(0, "/home/user/classifier")
import openpyxl
from openpyxl.styles import Font
import price_materials as PM

IN = "/home/user/classifier/ListaResurse.xlsx"
OUT = "/home/user/classifier/ListaResurse_preturi.xlsx"

wb = openpyxl.load_workbook(IN)
ws = wb["Materiale"]

# header for new columns (E already = "PRET UNITAR ")
hdr = Font(bold=True)
ws.cell(row=1, column=5, value="PRET UNITAR (RON, live)").font = hdr
ws.cell(row=1, column=6, value="Sursă").font = hdr
ws.cell(row=1, column=7, value="Produs potrivit (Dedeman)").font = hdr
ws.cell(row=1, column=8, value="Link sursă").font = hdr
ws.cell(row=1, column=9, value="Încredere potrivire").font = hdr
ws.column_dimensions['G'].width = 42
ws.column_dimensions['H'].width = 30
ws.column_dimensions['I'].width = 14

n_price = n_skip = n_none = 0
total = ws.max_row - 1
for i, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), start=2):
    name = row[2].value
    um = row[3].value
    if not name:
        continue
    try:
        res = PM.price_for(name, um)
    except Exception:
        res = {"price": None, "match_name": "", "match_url": "", "conf": "-",
               "note": "eroare"}
        traceback.print_exc()
    ws.cell(row=i, column=5, value=res["price"] if res["price"] else None)
    ws.cell(row=i, column=6, value=("Dedeman" if res["price"] else res.get("note", "")))
    ws.cell(row=i, column=7, value=res.get("match_name", "")[:120])
    lc = ws.cell(row=i, column=8, value=res.get("match_url", ""))
    if res.get("match_url"):
        lc.hyperlink = res["match_url"]
        lc.font = Font(color="1155CC", underline="single")
    ws.cell(row=i, column=9, value=res.get("conf", "-"))
    if res["price"]:
        n_price += 1
    elif res.get("note", "").startswith("poziție de tip"):
        n_skip += 1
    else:
        n_none += 1
    if i % 20 == 0:
        wb.save(OUT)
        print(f"[{i-1}/{total}] priced={n_price} skip={n_skip} none={n_none}", flush=True)
    time.sleep(0.8)

wb.save(OUT)
print(f"\nDONE. priced={n_price} allowance-skip={n_skip} none={n_none} of {total}", flush=True)
